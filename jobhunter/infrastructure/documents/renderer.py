"""Document rendering (Jinja CV template, optional WeasyPrint PDF) and CV text extraction."""
import io
from pathlib import Path

from jinja2 import Environment, FileSystemLoader, select_autoescape

from ...domain.errors import InvalidInput


class JinjaRenderer:
    def __init__(self, template_dir=None):
        self.env = Environment(loader=FileSystemLoader(template_dir or Path(__file__).parent / "templates"),
                               autoescape=select_autoescape(["html"]))

    def cv_html(self, context):
        return self.env.get_template("cv.html").render(**context)

    def pdf(self, html):
        """PDF bytes, or None if no renderer is available (the HTML is still produced).
        WeasyPrint first (Linux/Docker), then headless Chromium via Playwright (works on Windows)."""
        try:
            from weasyprint import HTML
            return HTML(string=html).write_pdf()
        except Exception:  # noqa: BLE001  not installed / missing GTK
            pass
        return _chromium_pdf(html)


def _chromium_pdf(html):
    """Print the HTML with headless Chromium: identical to the in-app preview. Uses Playwright's bundled
    browser (`python -m playwright install chromium`), else an installed Chrome or Edge."""
    try:
        from playwright.sync_api import sync_playwright
    except ImportError:
        return None
    try:
        with sync_playwright() as p:
            browser = None
            for opts in ({}, {"channel": "chrome"}, {"channel": "msedge"}):
                try:
                    browser = p.chromium.launch(**opts)
                    break
                except Exception:  # noqa: BLE001
                    continue
            if browser is None:
                return None
            try:
                page = browser.new_page()
                page.set_content(html, wait_until="load")
                return page.pdf(format="A4", print_background=True, prefer_css_page_size=True)
            finally:
                browser.close()
    except Exception:  # noqa: BLE001
        return None


class FileTextExtractor:
    supported = (".pdf", ".docx", ".md", ".markdown", ".txt", ".json", ".yaml", ".yml")

    def extract(self, filename, data):
        ext = ("." + filename.rsplit(".", 1)[-1].lower()) if "." in filename else ""
        if ext == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            return "\n".join(page.extract_text() or "" for page in reader.pages).strip()
        if ext == ".docx":
            from docx import Document
            doc = Document(io.BytesIO(data))
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
            return "\n".join(p for p in parts if p.strip())
        if ext in self.supported or not ext:
            return data.decode("utf-8", errors="replace")
        raise InvalidInput(f"Unsupported file type '{ext}'. Use PDF, DOCX, MD, TXT, JSON or YAML.")
